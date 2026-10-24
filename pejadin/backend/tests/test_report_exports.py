import asyncio
import io
import unittest
from unittest.mock import patch

from fastapi import HTTPException
from openpyxl import load_workbook
from pypdf import PdfReader

import main


class ReportExportTests(unittest.TestCase):
    def setUp(self):
        self.user = {"id": 1, "username": "operator", "role": "OPERATOR"}
        self.rows = [
            {
                "nomor": "ST/001/2026",
                "nama_kegiatan": "<Perjalanan> & tugas",
                "maksud": "Koordinasi",
                "asal": "Bandung",
                "provinsi_asal": "Jawa Barat",
                "tujuan": "Jakarta",
                "provinsi": "DKI Jakarta",
                "tgl_berangkat": "2026-10-01",
                "tgl_kembali": "2026-10-02",
                "tahun": 2026,
                "alat": "Kereta",
                "jml": 2,
                "nama_pegawai": ["Ayu", "=HYPERLINK(\"https://example.com\")"],
                "status": "TERBIT",
                "dibuat_oleh": "operator",
                "created_at": "2026-09-30T18:30:00+00:00",
                "diterbitkan_oleh": "ppk",
                "diterbitkan_pada": "2026-09-30",
                "total": 1250000,
                "flag": 1,
            },
            {
                "nomor": "ST/002/2025",
                "nama_kegiatan": "Data tahun sebelumnya",
                "tahun": 2025,
                "status": "DRAFT",
                "jml": 1,
                "total": 0,
            },
        ]
        self.satker = [
            {"k": "nama_kementerian", "v": "Kementerian Contoh"},
            {"k": "nama_ditjen", "v": "Direktorat Jenderal Contoh"},
            {"k": "nama_satker", "v": "Satuan Kerja Contoh"},
            {"k": "alamat", "v": "Jalan Merdeka"},
        ]

    def request_report(self, format_name, tahun=0, status=""):
        with patch.object(main, "stl", return_value=self.rows), patch.object(
            main, "q", return_value=self.satker
        ):
            return main.lap(format_name, tahun, status, self.user)

    def read_response(self, response):
        async def collect():
            return b"".join([chunk async for chunk in response.body_iterator])

        return asyncio.run(collect())

    def test_pdf_report_is_valid_formal_and_escapes_user_content(self):
        response = self.request_report("pdf", tahun=2026, status="TERBIT")
        self.assertEqual(response.media_type, "application/pdf")
        self.assertIn(
            "attachment; filename=\"laporan_rekap_surat_tugas_2026.pdf\"",
            response.headers["content-disposition"],
        )
        self.assertTrue(response.body.startswith(b"%PDF-"))

        reader = PdfReader(io.BytesIO(response.body))
        self.assertEqual(reader.metadata.title, "Laporan Rekapitulasi Surat Tugas")
        self.assertGreater(
            float(reader.pages[0].mediabox.width),
            float(reader.pages[0].mediabox.height),
        )
        report_text = " ".join(
            " ".join((page.extract_text() or "").split())
            for page in reader.pages
        )

        for expected_text in (
            "LAPORAN REKAPITULASI SURAT TUGAS",
            "RINCIAN SURAT TUGAS",
            "Jumlah pelaksana",
            "Akumulasi biaya",
            "Tahun Anggaran 2026",
            "<Perjalanan> & tugas",
            "1.250.000",
            "01-10-2026 01:30 WIB",
            "Tanggal & Waktu Dibuat (WIB)",
        ):
            self.assertIn(expected_text, report_text)
        self.assertNotIn("Data tahun sebelumnya", report_text)

        all_years_response = self.request_report("pdf")
        all_years_text = " ".join(
            " ".join((page.extract_text() or "").split())
            for page in PdfReader(io.BytesIO(all_years_response.body)).pages
        )
        self.assertIn("Tahun Anggaran 2025 sampai 2026", all_years_text)

    def test_pdf_report_with_no_matching_rows_still_produces_pdf(self):
        with patch.object(main, "stl", return_value=[]), patch.object(
            main, "q", return_value=self.satker
        ):
            response = main.lap("pdf", 2026, "TERBIT", self.user)

        self.assertEqual(response.media_type, "application/pdf")
        report_text = "\n".join(
            page.extract_text() or ""
            for page in PdfReader(io.BytesIO(response.body)).pages
        )
        self.assertIn("Tidak ada data Surat Tugas", report_text)
        self.assertIn("Jumlah Surat Tugas", report_text)

    def test_workbook_has_summary_and_filtered_detail_with_safe_text(self):
        response = self.request_report("xlsx", tahun=2026, status="TERBIT")
        workbook = load_workbook(io.BytesIO(self.read_response(response)), data_only=False)

        self.assertEqual(workbook.sheetnames, ["Ringkasan", "Rincian Surat Tugas"])
        summary = workbook["Ringkasan"]
        detail = workbook["Rincian Surat Tugas"]
        self.assertEqual(summary["A10"].value, 1)
        self.assertEqual(summary["E10"].value, 1250000)
        self.assertEqual(summary["B14"].value, 0)
        self.assertEqual(summary["B16"].value, 1)
        self.assertEqual(detail.max_row, 6)
        self.assertEqual(detail["C6"].value, "<Perjalanan> & tugas")
        self.assertEqual(detail["N6"].data_type, "s")
        self.assertEqual(detail["N6"].value, 'Ayu, =HYPERLINK("https://example.com")')
        self.assertEqual(detail["Q6"].value.isoformat(), "2026-10-01T01:30:00")
        self.assertEqual(detail["Q6"].number_format, "dd-mm-yyyy hh:mm")
        self.assertEqual(detail.auto_filter.ref, "A5:U6")
        self.assertEqual(detail.freeze_panes, "A6")

    def test_invalid_format_is_rejected(self):
        with patch.object(main, "stl", return_value=[]), patch.object(
            main, "q", return_value=[]
        ):
            with self.assertRaises(HTTPException) as error:
                main.lap("doc", 0, "", self.user)

        self.assertEqual(error.exception.status_code, 422)


if __name__ == "__main__":
    unittest.main()
